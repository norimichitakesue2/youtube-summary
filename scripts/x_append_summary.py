"""1日分のツイート集約サマリJSONを tweets_log.xlsx に追記する。

サマリJSONスキーマ（Claudeが data/tweet_summaries/<YYYY-MM-DD>.json に書く）:
{
  "date": "YYYY-MM-DD",
  "n_tweets": 543,
  "summary_3lines": "1行目\\n2行目\\n3行目",
  "key_topics": ["...", "..."],
  "positive_themes": ["..."],
  "negative_themes": ["..."],
  "hot_tweets": [
    {"text": "...", "url": "...", "likes": 1234, "author_handle": "...", "author_name": "..."}
  ],
  "summarized_at": "ISO8601"
}

使い方:
    python3 scripts/x_append_summary.py <YYYY-MM-DD>          # 1日指定
    python3 scripts/x_append_summary.py --batch                # 全日付を一括
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SUMMARIES_DIR = ROOT / "data" / "tweet_summaries"
SUMMARIES_DIR.mkdir(parents=True, exist_ok=True)
LOG_FILE = ROOT / "tweets_log.xlsx"

COLUMNS = [
    ("日付", 12),
    ("ツイート数", 10),
    ("3行サマリ", 70),
    ("ポジ要素", 50),
    ("ネガ要素", 50),
    ("キートピック", 40),
    ("注目ツイートTOP3", 80),
    ("要約日時", 22),
]


def ensure_workbook():
    if LOG_FILE.exists():
        wb = load_workbook(LOG_FILE)
        return wb, wb.active
    wb = Workbook()
    ws = wb.active
    ws.title = "tweets_log"
    fill = PatternFill("solid", fgColor="0F172A")
    font = Font(bold=True, color="FFFFFF")
    for i, (name, w) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:" + get_column_letter(len(COLUMNS)) + "1"
    return wb, ws


def fmt_list(xs):
    return "\n".join(f"・{x}" for x in xs) if xs else ""


def fmt_hot_tweets(hot):
    if not hot:
        return ""
    lines = []
    for h in hot[:3]:
        text = h.get("text", "")[:80]
        author = h.get("author_handle", "")
        likes = h.get("likes", 0)
        lines.append(f"@{author} (♥{likes}): {text}")
    return "\n\n".join(lines)


def already_logged(ws, date):
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True):
        if row[0] == date:
            return True
    return False


def append(summary):
    wb, ws = ensure_workbook()
    if already_logged(ws, summary["date"]):
        return False
    row = [
        summary.get("date", ""),
        summary.get("n_tweets", 0),
        summary.get("summary_3lines", ""),
        fmt_list(summary.get("positive_themes", [])),
        fmt_list(summary.get("negative_themes", [])),
        fmt_list(summary.get("key_topics", [])),
        fmt_hot_tweets(summary.get("hot_tweets", [])),
        summary.get("summarized_at", ""),
    ]
    ws.append(row)
    new_row = ws.max_row
    for i in range(1, len(COLUMNS) + 1):
        ws.cell(row=new_row, column=i).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    wb.save(LOG_FILE)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("date", nargs="?")
    ap.add_argument("--batch", action="store_true")
    args = ap.parse_args()
    if args.batch:
        added = skipped = 0
        for p in sorted(SUMMARIES_DIR.glob("*.json")):
            with p.open("r", encoding="utf-8") as f:
                s = json.load(f)
            if append(s):
                added += 1
            else:
                skipped += 1
        print(f"追記 {added} / 既存スキップ {skipped}")
    else:
        if not args.date:
            raise SystemExit("usage: x_append_summary.py <YYYY-MM-DD> | --batch")
        p = SUMMARIES_DIR / f"{args.date}.json"
        with p.open("r", encoding="utf-8") as f:
            s = json.load(f)
        ok = append(s)
        print(f"{'追記' if ok else 'スキップ(既存)'}: {args.date}")


if __name__ == "__main__":
    main()
