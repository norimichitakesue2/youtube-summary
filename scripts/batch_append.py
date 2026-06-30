"""data/pending_summaries/*.json のうち summary_3lines が入っているものを
video_log.xlsx に一括追記する。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import PENDING_DIR
from append_summary import COLUMNS, LOG_FILE, ensure_workbook, fmt_list


def main():
    # ワークブックは一度だけ読み込み・保存する（1ファイルごとに save すると
    # 件数が増えたとき非常に遅く、中断時にxlsxが壊れるため）。
    wb, ws = ensure_workbook()
    vid_col = COLUMNS.index(("video_id", 14)) + 1
    existing = set()
    for row in ws.iter_rows(min_row=2, min_col=vid_col, max_col=vid_col, values_only=True):
        if row[0]:
            existing.add(row[0])

    ncol = len(COLUMNS)
    files = sorted(PENDING_DIR.glob("*.json"))
    added = skipped = no_summary = 0
    for p in files:
        with p.open("r", encoding="utf-8") as f:
            d = json.load(f)
        if not d.get("summary_3lines"):
            no_summary += 1
            continue
        vid = d.get("video_id", "")
        if vid in existing:
            skipped += 1
            continue
        ws.append([
            d.get("published_at", "")[:10],
            d.get("channel_name", ""),
            d.get("title", ""),
            d.get("summary_3lines", ""),
            fmt_list(d.get("positive_points", [])),
            fmt_list(d.get("negative_points", [])),
            fmt_list(d.get("key_topics", [])),
            ", ".join(d.get("tags", [])),
            d.get("url", ""),
            vid,
            d.get("summarized_at", ""),
        ])
        r = ws.max_row
        for i in range(1, ncol + 1):
            ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical="top")
        existing.add(vid)
        added += 1

    wb.save(LOG_FILE)
    print(
        f"追記: {added} / 既存スキップ: {skipped} / 未要約: {no_summary} / 合計: {len(files)}"
    )


if __name__ == "__main__":
    main()
