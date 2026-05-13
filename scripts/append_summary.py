"""要約結果（JSON）を video_log.xlsx に1行追記する。

呼び出し方:
    python3 scripts/append_summary.py <summary.json>

summary.json のスキーマ:
{
  "video_id": "...",
  "channel_id": "...",
  "channel_name": "...",
  "tags": [...],
  "title": "...",
  "published_at": "ISO8601",
  "url": "...",
  "summary_3lines": "1行目\\n2行目\\n3行目",
  "key_topics": ["...", "..."],
  "positive_points": ["褒めてた点1", "..."],
  "negative_points": ["批判/不満点1", "..."],
  "summarized_at": "ISO8601"
}

このスクリプトは「メカニカルな追記」だけを担当。
要約本体（positive/negative 抽出）は Claude が pending_summaries/ の transcript を読んで生成する。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import LOG_FILE

COLUMNS = [
    ("公開日", 12),
    ("チャンネル", 18),
    ("タイトル", 60),
    ("3行要約", 70),
    ("ポジ要素（褒めてた点）", 50),
    ("ネガ要素（批判/不満）", 50),
    ("キートピック", 40),
    ("タグ", 20),
    ("URL", 45),
    ("video_id", 14),
    ("要約日時", 22),
]


def ensure_workbook() -> tuple:
    if LOG_FILE.exists():
        wb = load_workbook(LOG_FILE)
        ws = wb.active
        return wb, ws

    wb = Workbook()
    ws = wb.active
    ws.title = "video_log"
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")
    for i, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:" + get_column_letter(len(COLUMNS)) + "1"
    return wb, ws


def fmt_list(xs) -> str:
    if not xs:
        return ""
    return "\n".join(f"・{x}" for x in xs)


def already_logged(ws, video_id: str) -> bool:
    vid_col = COLUMNS.index(("video_id", 14)) + 1
    for row in ws.iter_rows(min_row=2, min_col=vid_col, max_col=vid_col, values_only=True):
        if row[0] == video_id:
            return True
    return False


def append(summary: dict) -> bool:
    wb, ws = ensure_workbook()
    if already_logged(ws, summary["video_id"]):
        wb.save(LOG_FILE)
        return False
    row = [
        summary.get("published_at", "")[:10],
        summary.get("channel_name", ""),
        summary.get("title", ""),
        summary.get("summary_3lines", ""),
        fmt_list(summary.get("positive_points", [])),
        fmt_list(summary.get("negative_points", [])),
        fmt_list(summary.get("key_topics", [])),
        ", ".join(summary.get("tags", [])),
        summary.get("url", ""),
        summary.get("video_id", ""),
        summary.get("summarized_at", ""),
    ]
    ws.append(row)
    new_row = ws.max_row
    for i in range(1, len(COLUMNS) + 1):
        ws.cell(row=new_row, column=i).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    wb.save(LOG_FILE)
    return True


def main():
    if len(sys.argv) < 2:
        raise SystemExit("usage: append_summary.py <summary.json>")
    path = Path(sys.argv[1])
    with path.open("r", encoding="utf-8") as f:
        summary = json.load(f)
    added = append(summary)
    print(f"{'追記' if added else 'スキップ(既存)'}: {summary['title'][:40]}")


if __name__ == "__main__":
    main()
