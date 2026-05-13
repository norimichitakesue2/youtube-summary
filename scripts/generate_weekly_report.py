"""video_log.xlsx を読み、直近1週間（または指定週）の動画をHTMLレポート化。

使い方:
    python3 scripts/generate_weekly_report.py            # 直近7日
    python3 scripts/generate_weekly_report.py --days 14
    python3 scripts/generate_weekly_report.py --since 2026-04-01 --until 2026-04-30

出力:
    reports/weekly_report_<until_date>.html
    reports/index.html  ← 既存レポートへのリンク一覧
"""
from __future__ import annotations

import argparse
import html
import json
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import LOG_FILE, PENDING_DIR, REPORTS_DIR

ROOT = Path(__file__).resolve().parent.parent
TWEET_SUMMARIES_DIR = ROOT / "data" / "tweet_summaries"

X_CHIP_LABEL = "X (Twitter)"


HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="ja">
<head>
<meta charset="utf-8">
<title>YouTube要約レポート {since} 〜 {until}</title>
<style>
  :root {{
    --bg: #0f172a;
    --panel: #1e293b;
    --panel-2: #273449;
    --text: #e2e8f0;
    --muted: #94a3b8;
    --accent: #60a5fa;
    --pos: #34d399;
    --neg: #f87171;
    --topic: #fbbf24;
  }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "Hiragino Kaku Gothic ProN", "Yu Gothic", sans-serif;
    background: var(--bg);
    color: var(--text);
    line-height: 1.6;
  }}
  header {{
    padding: 32px 48px;
    background: linear-gradient(135deg, #1e3a8a 0%, #312e81 100%);
    border-bottom: 1px solid #334155;
  }}
  header h1 {{ margin: 0 0 8px; font-size: 28px; }}
  header .meta {{ color: #cbd5e1; font-size: 14px; }}
  main {{ padding: 32px 48px; max-width: 1200px; margin: 0 auto; }}
  .summary-stats {{
    display: flex;
    gap: 16px;
    margin-bottom: 32px;
    flex-wrap: wrap;
  }}
  .stat {{
    background: var(--panel);
    padding: 16px 24px;
    border-radius: 8px;
    flex: 1;
    min-width: 180px;
  }}
  .stat .num {{ font-size: 28px; font-weight: 700; color: var(--accent); }}
  .stat .label {{ font-size: 13px; color: var(--muted); margin-top: 4px; }}
  .date-section {{ margin-bottom: 32px; }}
  .date-section h2 {{
    color: var(--accent);
    border-bottom: 2px solid var(--accent);
    padding-bottom: 8px;
    margin-bottom: 16px;
    font-size: 20px;
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    gap: 12px;
  }}
  .date-section h2 .day-count {{
    font-size: 13px;
    font-weight: normal;
    color: var(--muted);
  }}
  .video {{
    background: var(--panel);
    border-radius: 10px;
    padding: 20px 24px;
    margin-bottom: 12px;
    border-left: 4px solid var(--accent);
  }}
  .video-head {{
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    gap: 16px;
    margin-bottom: 12px;
    flex-wrap: wrap;
  }}
  .video-channel {{
    display: inline-block;
    background: rgba(96, 165, 250, 0.18);
    color: var(--accent);
    padding: 2px 10px;
    border-radius: 10px;
    font-size: 12px;
    font-weight: 500;
    margin-bottom: 6px;
  }}
  .video.pending {{ border-left: 3px solid var(--muted); opacity: 0.85; }}
  .pending-badge {{
    display: inline-block;
    background: rgba(148, 163, 184, 0.18);
    color: var(--muted);
    padding: 2px 10px;
    border-radius: 10px;
    font-size: 11px;
    font-weight: 500;
    margin-left: 6px;
    margin-bottom: 6px;
  }}
  .video-title {{
    font-size: 16px;
    font-weight: 600;
    margin: 0;
  }}
  .video-title a {{ color: var(--text); text-decoration: none; }}
  .video-title a:hover {{ color: var(--accent); }}
  .video-date {{ color: var(--muted); font-size: 13px; white-space: nowrap; }}
  .video-summary {{
    color: #cbd5e1;
    margin: 8px 0 16px;
    white-space: pre-wrap;
  }}
  .pos-neg-grid {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 12px;
    margin-bottom: 12px;
  }}
  @media (max-width: 700px) {{ .pos-neg-grid {{ grid-template-columns: 1fr; }} }}
  .pos, .neg {{
    background: var(--panel-2);
    border-radius: 6px;
    padding: 10px 14px;
    font-size: 14px;
  }}
  .pos {{ border-left: 3px solid var(--pos); }}
  .neg {{ border-left: 3px solid var(--neg); }}
  .pos .label, .neg .label {{
    font-size: 12px;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-bottom: 6px;
  }}
  .pos .label {{ color: var(--pos); }}
  .neg .label {{ color: var(--neg); }}
  .pos ul, .neg ul {{ margin: 0; padding-left: 18px; }}
  .pos li, .neg li {{ margin-bottom: 2px; }}
  .topics {{ margin-top: 8px; display: flex; flex-wrap: wrap; gap: 6px; }}
  .topic {{
    background: rgba(251, 191, 36, 0.15);
    color: var(--topic);
    padding: 3px 10px;
    border-radius: 12px;
    font-size: 12px;
  }}
  .empty-note {{ color: var(--muted); font-size: 13px; font-style: italic; }}
  .filter-bar {{
    background: var(--panel);
    border-radius: 10px;
    padding: 16px 20px;
    margin-bottom: 24px;
    position: sticky;
    top: 0;
    z-index: 10;
    box-shadow: 0 4px 12px rgba(0,0,0,0.3);
  }}
  .filter-bar .filter-label {{
    font-size: 12px;
    color: var(--muted);
    margin-bottom: 8px;
    display: flex;
    justify-content: space-between;
    align-items: center;
  }}
  .filter-bar .filter-actions {{
    display: flex;
    gap: 8px;
  }}
  .filter-bar button {{
    background: transparent;
    color: var(--muted);
    border: 1px solid #334155;
    padding: 3px 10px;
    border-radius: 12px;
    font-size: 11px;
    cursor: pointer;
    transition: all 0.15s;
  }}
  .filter-bar button:hover {{
    color: var(--accent);
    border-color: var(--accent);
  }}
  .filter-bar .chips {{
    display: flex;
    flex-wrap: wrap;
    gap: 8px;
  }}
  .filter-bar .chip {{
    background: var(--panel-2);
    color: var(--muted);
    padding: 6px 14px;
    border-radius: 16px;
    font-size: 13px;
    cursor: pointer;
    user-select: none;
    border: 1px solid transparent;
    transition: all 0.15s;
  }}
  .filter-bar .chip:hover {{
    border-color: var(--accent);
  }}
  .filter-bar .chip.active {{
    background: var(--accent);
    color: #0f172a;
    font-weight: 600;
  }}
  .filter-bar .chip.x-chip {{
    border-color: #1da1f2;
    color: #7dd3fc;
  }}
  .filter-bar .chip.x-chip.active {{
    background: #1da1f2;
    color: #0f172a;
  }}
  .filter-bar .chip .count {{
    opacity: 0.7;
    font-size: 11px;
    margin-left: 4px;
  }}
  .video.hidden {{ display: none; }}
  .date-section.hidden {{ display: none; }}
  .x-card.hidden {{ display: none; }}
  .x-card {{
    background: linear-gradient(180deg, #0c4a6e 0%, #164e63 100%);
    border-radius: 10px;
    padding: 18px 22px;
    margin-bottom: 12px;
    border-left: 4px solid #1da1f2;
  }}
  .x-badge {{
    display: inline-block;
    background: rgba(29, 161, 242, 0.25);
    color: #7dd3fc;
    padding: 2px 10px;
    border-radius: 10px;
    font-size: 12px;
    font-weight: 600;
    margin-bottom: 6px;
  }}
  .x-card h3 {{
    margin: 0 0 8px;
    font-size: 16px;
    font-weight: 600;
  }}
  .x-card .x-summary {{
    color: #cbd5e1;
    margin-bottom: 14px;
    white-space: pre-wrap;
    font-size: 14px;
  }}
  .x-card .hot {{
    margin-top: 12px;
  }}
  .x-card .hot-label {{
    color: #94a3b8;
    font-size: 11px;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-bottom: 6px;
  }}
  .x-card .tweet {{
    background: rgba(15, 23, 42, 0.55);
    border-radius: 6px;
    padding: 10px 14px;
    margin-bottom: 6px;
    font-size: 13px;
  }}
  .x-card .tweet .author {{
    color: #7dd3fc;
    font-weight: 600;
    margin-bottom: 2px;
    font-size: 12px;
  }}
  .x-card .tweet a {{
    color: #7dd3fc;
    text-decoration: none;
    font-size: 11px;
  }}
  .x-card .tweet .stats {{
    color: #94a3b8;
    font-size: 11px;
    margin-top: 4px;
  }}
  .visible-stats {{
    display: inline-block;
    color: var(--muted);
    font-size: 12px;
    margin-left: 8px;
  }}
  footer {{
    text-align: center;
    color: var(--muted);
    font-size: 12px;
    padding: 24px;
    border-top: 1px solid #334155;
  }}
</style>
</head>
<body>
<header>
  <h1>YouTube要約レポート</h1>
  <div class="meta">対象期間: {since} 〜 {until}（{ndays}日間） / 生成日時: {generated}</div>
</header>
<main>
  <div class="summary-stats">
    <div class="stat"><div class="num">{n_videos}</div><div class="label">対象動画数</div></div>
    <div class="stat"><div class="num">{n_channels}</div><div class="label">チャンネル数</div></div>
    <div class="stat"><div class="num">{n_pos}</div><div class="label">ポジ要素（合計）</div></div>
    <div class="stat"><div class="num">{n_neg}</div><div class="label">ネガ要素（合計）</div></div>
  </div>
  <div class="filter-bar" id="filterBar">
    <div class="filter-label">
      <span>表示するYouTuberを選択 <span id="visibleStats" class="visible-stats"></span></span>
      <div class="filter-actions">
        <button onclick="setAll(true)">全て表示</button>
        <button onclick="setAll(false)">全て非表示</button>
      </div>
    </div>
    <div class="chips" id="chips">{chips}</div>
  </div>
  {body}
</main>
<footer>generated by youtube-summary</footer>
<script>
function getActiveChannels() {{
  return new Set(
    Array.from(document.querySelectorAll('.chip.active'))
      .map(c => c.dataset.channel)
  );
}}
function applyFilter() {{
  const active = getActiveChannels();
  // 各動画の表示/非表示
  document.querySelectorAll('.video').forEach(v => {{
    const ch = v.dataset.channel;
    v.classList.toggle('hidden', !active.has(ch));
  }});
  // Xカードの表示/非表示
  document.querySelectorAll('.x-card').forEach(v => {{
    const ch = v.dataset.channel;
    v.classList.toggle('hidden', !active.has(ch));
  }});
  // 日付セクションは、コンテンツが0なら非表示
  document.querySelectorAll('.date-section').forEach(s => {{
    const visible = s.querySelectorAll('.video:not(.hidden), .x-card:not(.hidden)').length;
    s.classList.toggle('hidden', visible === 0);
  }});
  updateStats();
}}
function updateStats() {{
  const totalCh = document.querySelectorAll('.chip').length;
  const activeCh = document.querySelectorAll('.chip.active').length;
  const visibleVideos = document.querySelectorAll('.video:not(.hidden)').length;
  const visibleX = document.querySelectorAll('.x-card:not(.hidden)').length;
  const visibleDays = document.querySelectorAll('.date-section:not(.hidden)').length;
  document.getElementById('visibleStats').textContent =
    `（${{activeCh}}/${{totalCh}} 選択中、動画${{visibleVideos}}本／X${{visibleX}}日／合計${{visibleDays}}日）`;
}}
function toggleChannel(chip) {{
  chip.classList.toggle('active');
  applyFilter();
}}
function setAll(on) {{
  document.querySelectorAll('.chip').forEach(c => c.classList.toggle('active', on));
  applyFilter();
}}
document.querySelectorAll('.chip').forEach(c => {{
  c.addEventListener('click', () => toggleChannel(c));
}});
updateStats();
</script>
</body>
</html>
"""


def parse_list_cell(s) -> list[str]:
    if not s:
        return []
    return [line.lstrip("・").strip() for line in str(s).split("\n") if line.strip()]


def render_video(v: dict) -> str:
    if v.get("pending"):
        return render_pending_video(v)
    title = html.escape(v["title"])
    url = html.escape(v["url"])
    channel = html.escape(v["channel_name"])
    channel_attr = html.escape(v["channel_name"], quote=True)
    summary = html.escape(v.get("summary_3lines", "")).replace("\n", "<br>")
    pos = v.get("positive_points", [])
    neg = v.get("negative_points", [])
    topics = v.get("key_topics", [])

    pos_html = (
        "<ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in pos) + "</ul>"
        if pos
        else '<div class="empty-note">特になし</div>'
    )
    neg_html = (
        "<ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in neg) + "</ul>"
        if neg
        else '<div class="empty-note">特になし</div>'
    )
    topics_html = "".join(
        f'<span class="topic">{html.escape(t)}</span>' for t in topics
    )

    return f"""
<div class="video" data-channel="{channel_attr}">
  <div>
    <span class="video-channel">{channel}</span>
  </div>
  <div class="video-head">
    <h3 class="video-title"><a href="{url}" target="_blank" rel="noopener">{title}</a></h3>
  </div>
  <div class="video-summary">{summary or '<span class="empty-note">要約なし</span>'}</div>
  <div class="pos-neg-grid">
    <div class="pos"><div class="label">ポジ要素</div>{pos_html}</div>
    <div class="neg"><div class="label">ネガ要素</div>{neg_html}</div>
  </div>
  {f'<div class="topics">{topics_html}</div>' if topics_html else ''}
</div>
"""


def render_pending_video(v: dict) -> str:
    title = html.escape(v["title"])
    url = html.escape(v["url"])
    channel = html.escape(v["channel_name"])
    channel_attr = html.escape(v["channel_name"], quote=True)
    status = v.get("transcript_status", "")
    if status == "no_transcript":
        badge_text = "字幕なし・未要約"
    elif "IpBlocked" in status or "RequestBlocked" in status:
        badge_text = "字幕取得待ち（IPブロック中）"
    elif status.startswith("error"):
        badge_text = "字幕取得エラー・未要約"
    elif status == "ok":
        badge_text = "要約待ち"
    else:
        badge_text = "未要約"
    return f"""
<div class="video pending" data-channel="{channel_attr}">
  <div>
    <span class="video-channel">{channel}</span>
    <span class="pending-badge">{html.escape(badge_text)}</span>
  </div>
  <div class="video-head">
    <h3 class="video-title"><a href="{url}" target="_blank" rel="noopener">{title}</a></h3>
  </div>
</div>
"""


def render_x_card(s: dict) -> str:
    summary = html.escape(s.get("summary_3lines", "")).replace("\n", "<br>")
    pos = s.get("positive_themes", [])
    neg = s.get("negative_themes", [])
    topics = s.get("key_topics", [])
    hot = s.get("hot_tweets", [])

    pos_html = (
        "<ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in pos) + "</ul>"
        if pos
        else '<div class="empty-note">特になし</div>'
    )
    neg_html = (
        "<ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in neg) + "</ul>"
        if neg
        else '<div class="empty-note">特になし</div>'
    )
    topics_html = "".join(
        f'<span class="topic">{html.escape(t)}</span>' for t in topics
    )
    hot_html = ""
    for t in hot[:3]:
        text = html.escape(t.get("text", "")).replace("\n", "<br>")
        author = html.escape(t.get("author_handle", ""))
        url = html.escape(t.get("url", "#"))
        likes = t.get("likes", 0)
        rts = t.get("retweets", 0)
        rep = t.get("replies", 0)
        hot_html += f"""
<div class="tweet">
  <div class="author">@{author}</div>
  <div>{text}</div>
  <div class="stats">♥{likes} ↻{rts} 💬{rep} ・ <a href="{url}" target="_blank" rel="noopener">原文</a></div>
</div>"""

    n_tweets = s.get("n_tweets", 0)
    return f"""
<div class="x-card" data-channel="{X_CHIP_LABEL}">
  <div><span class="x-badge">X (Twitter) ・ {n_tweets}ツイート集約</span></div>
  <h3>その日のXハイライト</h3>
  <div class="x-summary">{summary or '<span class="empty-note">サマリなし</span>'}</div>
  <div class="pos-neg-grid">
    <div class="pos"><div class="label">ポジ要素（X）</div>{pos_html}</div>
    <div class="neg"><div class="label">ネガ要素（X）</div>{neg_html}</div>
  </div>
  {f'<div class="topics">{topics_html}</div>' if topics_html else ''}
  {f'<div class="hot"><div class="hot-label">注目ツイート TOP{min(3,len(hot))}</div>{hot_html}</div>' if hot else ''}
</div>
"""


def load_x_summaries(since: datetime, until: datetime) -> dict[str, dict]:
    out: dict[str, dict] = {}
    if not TWEET_SUMMARIES_DIR.exists():
        return out
    for p in TWEET_SUMMARIES_DIR.glob("*.json"):
        try:
            d = datetime.fromisoformat(p.stem)
        except ValueError:
            continue
        if not (since <= d <= until):
            continue
        with p.open("r", encoding="utf-8") as f:
            out[p.stem] = json.load(f)
    return out


def load_videos(since: datetime, until: datetime) -> list[dict]:
    if not LOG_FILE.exists():
        return []
    wb = load_workbook(LOG_FILE)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    videos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        date_str = row[idx["公開日"]]
        if not date_str:
            continue
        try:
            d = datetime.fromisoformat(str(date_str)[:10])
        except ValueError:
            continue
        if not (since <= d <= until):
            continue
        videos.append(
            {
                "published_at": str(date_str)[:10],
                "channel_name": row[idx["チャンネル"]] or "",
                "title": row[idx["タイトル"]] or "",
                "summary_3lines": row[idx["3行要約"]] or "",
                "positive_points": parse_list_cell(row[idx["ポジ要素（褒めてた点）"]]),
                "negative_points": parse_list_cell(row[idx["ネガ要素（批判/不満）"]]),
                "key_topics": parse_list_cell(row[idx["キートピック"]]),
                "url": row[idx["URL"]] or "",
            }
        )
    videos.sort(key=lambda v: (v["channel_name"], v["published_at"]), reverse=False)
    return videos


def load_pending_videos(
    since: datetime, until: datetime, exclude_urls: set[str]
) -> list[dict]:
    """pending_summaries/ から、要約未生成だが期間内の動画を読み込む。"""
    if not PENDING_DIR.exists():
        return []
    pending = []
    for p in sorted(PENDING_DIR.glob("*.json")):
        if p.name.startswith("_"):
            continue
        try:
            with p.open("r", encoding="utf-8") as f:
                d = json.load(f)
        except Exception:
            continue
        if d.get("summary_3lines"):
            continue
        url = d.get("url", "")
        if url and url in exclude_urls:
            continue
        pub = d.get("published_at", "")
        if not pub:
            continue
        try:
            dt = datetime.fromisoformat(pub.replace("Z", "+00:00")).replace(
                tzinfo=None
            )
        except ValueError:
            continue
        if not (since <= dt <= until + timedelta(days=1)):
            continue
        pending.append(
            {
                "published_at": pub[:10],
                "channel_name": d.get("channel_name", ""),
                "title": d.get("title", ""),
                "url": url,
                "transcript_status": d.get("transcript_status", ""),
                "summary_3lines": "",
                "positive_points": [],
                "negative_points": [],
                "key_topics": [],
                "pending": True,
            }
        )
    return pending


def render_index() -> None:
    import re

    files = sorted(REPORTS_DIR.glob("weekly_report_*.html"), reverse=True)
    JP_WD = ["月", "火", "水", "木", "金", "土", "日"]

    # ファイル名から until 日付を抽出してメタを作る
    entries = []  # list of dict(name, until_dt, since_dt, is_monday)
    pat = re.compile(r"weekly_report_(\d{4})-(\d{2})-(\d{2})\.html")
    for f in files:
        m = pat.match(f.name)
        if not m:
            continue
        until = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        since = until - timedelta(days=6)
        entries.append(
            {
                "name": f.name,
                "until": until,
                "since": since,
                "is_monday": until.weekday() == 0,
            }
        )

    # 月別グルーピング
    by_month: dict[str, list[dict]] = defaultdict(list)
    for e in entries:
        by_month[e["until"].strftime("%Y-%m")].append(e)
    months_sorted = sorted(by_month.keys(), reverse=True)

    # X レポートも併記（存在すれば）
    x_index_link = ""
    if (REPORTS_DIR / "x_index.html").exists():
        x_index_link = (
            '<p class="ext-link">'
            '🐦 X (Twitter) サマリ一覧: '
            '<a href="x_index.html">x_index.html</a></p>'
        )

    # 最新ハイライト（本日分）
    latest = entries[0] if entries else None
    if latest:
        latest_html = (
            f'<div class="latest">'
            f'<div class="latest-label">最新（本日の更新分）</div>'
            f'<a class="latest-link" href="{latest["name"]}">'
            f'{latest["until"].strftime("%Y-%m-%d")} '
            f'<span class="range">({latest["since"].strftime("%-m/%-d")}〜'
            f'{latest["until"].strftime("%-m/%-d")} の7日間)</span>'
            f'</a></div>'
        )
    else:
        latest_html = ""

    # 週次まとめ（weekly_summary_*.html を最上段に大きく表示）
    weekly_summary_files = sorted(
        REPORTS_DIR.glob("weekly_summary_*.html"), reverse=True
    )
    weekly_summary_entries = []
    pat_ws = re.compile(r"weekly_summary_(\d{4})-(\d{2})-(\d{2})\.html")
    for f in weekly_summary_files:
        m = pat_ws.match(f.name)
        if not m:
            continue
        until = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        since = until - timedelta(days=6)
        weekly_summary_entries.append({"name": f.name, "until": until, "since": since})

    weekly_html = ""
    if weekly_summary_entries:
        cards = []
        for e in weekly_summary_entries:
            cards.append(
                f'<a class="weekly-card" href="{e["name"]}">'
                f'<div class="weekly-card-date">{e["until"].strftime("%Y-%m-%d")} 週次総括</div>'
                f'<div class="weekly-card-range">{e["since"].strftime("%-m/%-d")}〜'
                f'{e["until"].strftime("%-m/%-d")}</div>'
                f'</a>'
            )
        weekly_html = (
            f'<section class="weekly-section">'
            f'<h2>📌 週次総括（メタサマリ） <span class="count">({len(weekly_summary_entries)}件)</span></h2>'
            f'<p class="section-desc">毎週月曜の <code>youtube-summary-weekly</code> '
            f'タスクで生成される、1週間分の動画群を <b>横断的に集約・要約した上位レイヤー</b>のレポート。'
            f'話題まとめ／頻出トピック／チャンネル別ハイライト／注目ピックアップを掲載。</p>'
            f'<div class="weekly-grid">{"".join(cards)}</div>'
            f'</section>'
        )

    # 日次スナップショットは月別アーカイブ（weekly_summaryが既に作られた日付は除外しない＝
    # 月曜の日次スナップショットも一覧に残す）
    section_parts = []
    for ym in months_sorted:
        items = by_month[ym]
        if not items:
            continue
        y, m = ym.split("-")
        section_parts.append(
            f'<section class="month"><h2>📅 {y}年{int(m)}月 '
            f'<span class="count">({len(items)}件)</span></h2><ul>'
        )
        for e in items:
            wd = JP_WD[e["until"].weekday()]
            section_parts.append(
                f'<li>'
                f'<a href="{e["name"]}">{e["until"].strftime("%Y-%m-%d")}（{wd}）</a>'
                f' <span class="range">{e["since"].strftime("%-m/%-d")}〜'
                f'{e["until"].strftime("%-m/%-d")}</span>'
                f'</li>'
            )
        section_parts.append("</ul></section>")

    archive_wrap = ""
    if section_parts:
        archive_wrap = (
            f'<h2 class="archive-h2">🗂 日次スナップショット（アーカイブ）</h2>'
            f'<p class="section-desc">日次タスクが毎日生成している、その日までの過去7日間レポート。</p>'
            f'{"".join(section_parts)}'
        )

    note = ""  # 説明は各セクションのsection-descに移動

    css = """body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI","Hiragino Kaku Gothic ProN","Yu Gothic",sans-serif;max-width:820px;margin:32px auto;padding:0 16px;color:#1f2937;line-height:1.6}
h1{font-size:24px;margin-bottom:4px}
.subtitle{color:#6b7280;font-size:14px;margin-top:0}
a{color:#2563eb;text-decoration:none}
a:hover{text-decoration:underline}
.section-desc{color:#6b7280;font-size:13px;margin:4px 0 12px}
.section-desc code{background:#e5e7eb;padding:1px 6px;border-radius:3px;font-size:12px}
.latest{background:linear-gradient(135deg,#dbeafe 0%,#e0e7ff 100%);padding:16px 20px;border-radius:10px;margin:20px 0 24px;border-left:4px solid #2563eb}
.latest-label{font-size:12px;font-weight:700;color:#1d4ed8;letter-spacing:0.05em;margin-bottom:4px}
.latest-link{font-size:18px;font-weight:600}
.latest-link .range{font-size:13px;font-weight:400;color:#4b5563}
.ext-link{font-size:14px;color:#4b5563;margin:0 0 16px}
.weekly-section{background:#fffbeb;border:1px solid #fde68a;border-radius:10px;padding:18px 20px;margin:24px 0}
.weekly-section h2{font-size:18px;margin:0 0 4px;color:#92400e}
.weekly-section h2 .count{font-size:12px;font-weight:400;color:#a16207;margin-left:6px}
.weekly-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(170px,1fr));gap:10px;margin-top:8px}
.weekly-card{display:block;background:#fff;border:1px solid #fde68a;border-radius:8px;padding:10px 12px;text-decoration:none;color:#92400e;transition:all 0.15s}
.weekly-card:hover{background:#fef3c7;border-color:#f59e0b;text-decoration:none;transform:translateY(-1px)}
.weekly-card-date{font-weight:600;font-size:14px}
.weekly-card-range{color:#a16207;font-size:12px;margin-top:2px}
.archive-h2{font-size:18px;margin:32px 0 4px;color:#374151;border-bottom:1px solid #e5e7eb;padding-bottom:6px}
.month{margin-bottom:20px}
.month h2{font-size:15px;margin:20px 0 6px;color:#4b5563;font-weight:600}
.month h2 .count{font-size:12px;font-weight:400;color:#9ca3af;margin-left:6px}
.month ul{list-style:none;padding-left:0;margin:0}
.month li{padding:3px 0;font-size:13px;color:#6b7280}
.month li a{color:#2563eb}
.month li .range{color:#9ca3af;font-size:12px;margin-left:6px}"""

    body = (
        f'<h1>YouTube要約レポート一覧</h1>'
        f'<p class="subtitle">日次タスクで毎日更新（最新7日分のスナップショット）</p>'
        f'{latest_html}'
        f'{x_index_link}'
        f'{weekly_html}'
        f'{archive_wrap}'
    )

    (REPORTS_DIR / "index.html").write_text(
        f'<!DOCTYPE html>\n<html lang="ja"><head><meta charset="utf-8">'
        f'<title>YouTube要約レポート一覧</title>'
        f'<style>{css}</style></head>\n<body>{body}</body></html>',
        encoding="utf-8",
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--days", type=int, default=7)
    ap.add_argument("--since")
    ap.add_argument("--until")
    args = ap.parse_args()

    if args.until:
        until = datetime.fromisoformat(args.until)
    else:
        until = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    if args.since:
        since = datetime.fromisoformat(args.since)
    else:
        since = until - timedelta(days=args.days - 1)

    videos = load_videos(since, until)
    existing_urls = {v["url"] for v in videos if v.get("url")}
    pending_videos = load_pending_videos(since, until, existing_urls)
    x_summaries = load_x_summaries(since, until)

    # 日付ごとにグループ化（新しい順、YouTube + X両方の日付を含む）
    by_date_videos = defaultdict(list)
    for v in videos:
        by_date_videos[v["published_at"]].append(v)
    for v in pending_videos:
        by_date_videos[v["published_at"]].append(v)
    all_dates = sorted(
        set(by_date_videos.keys()) | set(x_summaries.keys()), reverse=True
    )

    # チャンネルごとの本数（チップ用、動画数の多い順）
    by_channel = defaultdict(list)
    for v in videos:
        by_channel[v["channel_name"]].append(v)
    for v in pending_videos:
        by_channel[v["channel_name"]].append(v)
    sorted_channels = sorted(by_channel.items(), key=lambda x: -len(x[1]))

    body_parts = []
    for date in all_dates:
        vids = by_date_videos.get(date, [])
        x_sum = x_summaries.get(date)
        # 日付内ではチャンネル名→タイトルでソート
        vids_sorted = sorted(vids, key=lambda x: (x["channel_name"], x["title"]))
        contents_html = "".join(render_video(v) for v in vids_sorted)
        if x_sum:
            contents_html += render_x_card(x_sum)
        n_videos = len(vids)
        n_tweets = x_sum.get("n_tweets", 0) if x_sum else 0
        count_label = f"{n_videos}本" + (f" ／ X {n_tweets}ツイート" if x_sum else "")
        body_parts.append(
            f'<section class="date-section">'
            f'<h2>{html.escape(date)}<span class="day-count">{count_label}</span></h2>'
            f'{contents_html}</section>'
        )

    chips_parts = []
    for ch, vids in sorted_channels:
        ch_esc = html.escape(ch)
        ch_attr = html.escape(ch, quote=True)
        chips_parts.append(
            f'<span class="chip active" data-channel="{ch_attr}">{ch_esc}'
            f'<span class="count">{len(vids)}</span></span>'
        )
    # Xのチップ（サマリがある場合のみ）
    if x_summaries:
        chips_parts.append(
            f'<span class="chip x-chip active" data-channel="{html.escape(X_CHIP_LABEL, quote=True)}">'
            f'{html.escape(X_CHIP_LABEL)}<span class="count">{len(x_summaries)}日</span></span>'
        )

    if not body_parts:
        body_parts.append(
            '<p style="color:#94a3b8">この期間に対象データはありませんでした。</p>'
        )

    n_pos = sum(len(v["positive_points"]) for v in videos)
    n_neg = sum(len(v["negative_points"]) for v in videos)

    html_out = HTML_TEMPLATE.format(
        since=since.strftime("%Y-%m-%d"),
        until=until.strftime("%Y-%m-%d"),
        ndays=(until - since).days + 1,
        generated=datetime.now().strftime("%Y-%m-%d %H:%M"),
        n_videos=len(videos),
        n_channels=len(by_channel),
        n_pos=n_pos,
        n_neg=n_neg,
        body="\n".join(body_parts),
        chips="\n".join(chips_parts) if videos else "",
    )

    out = REPORTS_DIR / f"weekly_report_{until.strftime('%Y-%m-%d')}.html"
    out.write_text(html_out, encoding="utf-8")
    render_index()
    print(f"レポート生成: {out}")
    print(f"  対象期間: {since.date()} 〜 {until.date()}")
    print(
        f"  動画数: {len(videos)} (要約済) + {len(pending_videos)} (要約待ち) "
        f"/ チャンネル数: {len(by_channel)}"
    )


if __name__ == "__main__":
    main()
